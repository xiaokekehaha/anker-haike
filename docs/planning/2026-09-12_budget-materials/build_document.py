from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / 'docs/Anker黑客松预赛材料整理与填写模板.md'
OUTPUT = SOURCE.with_suffix('.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.75), Inches(.75)

for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'List Bullet']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
    fonts = st.element.get_or_add_rPr().rFonts
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower(): del fonts.attrib[attr]
    fonts.set(qn('w:eastAsia'), 'Heiti SC')
    for border in st.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.12
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(12)
for name, size in [('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 12), ('Heading 4', 11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(5)

def inline(p, text):
    tokens = re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\))', text)
    for t in tokens:
        if not t:
            continue
        if t.startswith('**'):
            p.add_run(t[2:-2]).bold = True
        elif re.match(r'\[[^\]]+\]\(https?://', t):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', t)
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), p.part.relate_to(m.group(2), RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            rp = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '205B86'); rp.append(color)
            run.append(rp)
            tx = OxmlElement('w:t'); tx.text = m.group(1); run.append(tx)
            link.append(run); p._p.append(link)
        else:
            p.add_run(t)

def table(rows):
    count = len(rows[0])
    tbl = doc.add_table(rows=1, cols=count)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    widths = [1.2, 5.8] if count == 2 else [1.25, 1.8, 1.85, 2.1]
    for col, width in zip(tbl.columns, widths): col.width = Inches(width)
    for idx, row in enumerate(rows):
        cells = tbl.rows[0].cells if idx == 0 else tbl.add_row().cells
        for j, val in enumerate(row):
            cells[j].width = Inches(widths[j]); cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.08
            inline(p, val)
            for r in p.runs: r.font.size = Pt(10)
            tcpr = cells[j]._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DCE8F0' if idx == 0 else ('F5F7F9' if idx % 2 == 0 else 'FFFFFF')); tcpr.append(sh)
            margins = OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                elem = OxmlElement(f'w:{edge}'); elem.set(qn('w:w'), '85'); elem.set(qn('w:type'), 'dxa'); margins.append(elem)
            tcpr.append(margins)
            if idx == 0:
                for r in p.runs: r.bold = True
        pr = tbl.rows[idx]._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if idx == 0: pr.append(OxmlElement('w:tblHeader'))
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x = OxmlElement('w:'+edge); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), 'D9D9D9'); borders.append(x)
    tbl._tbl.tblPr.append(borders)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = Pt(3)
    p.add_run().font.size = Pt(3)

lines = SOURCE.read_text().splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line == '---':
        doc.add_page_break(); i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', v) for v in cells): rows.append(cells)
            i += 1
        table(rows); continue
    if line.startswith('#'):
        n = len(line)-len(line.lstrip('#'))
        content = line[n:].strip()
        p = doc.add_paragraph(style='Title' if n == 1 else f'Heading {n-1}')
        inline(p, content)
    else:
        p = doc.add_paragraph(style='List Bullet' if line.startswith('- ') else 'Normal')
        inline(p, line[2:] if line.startswith('- ') else line)
    i += 1

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run(); run.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); run._r.addnext(field)
doc.core_properties.title = 'Anker首届黑客松挑战赛 预赛材料整理与填写模板'
doc.core_properties.subject = '预赛材料要求与可填写结构'
doc.core_properties.author = ''
doc.core_properties.keywords = 'Anker 黑客松 预赛 材料'
doc.save(OUTPUT)
print(OUTPUT)
