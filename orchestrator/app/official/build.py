"""把主办方 xlsx 导出为 data/official/*.yaml 与附件缩略图（确定性，可重复执行）。

    uv run python -m app.official.build [--source XLSX] [--out DIR]

只做结构化与归一化，不改写业务含义：
- 订单 sheet：带订单号的行开启一个订单，其后无订单号的行是该订单的子行（配件或拆单）；
  多包裹单元格按换行对齐拆成 Package。
- 客诉 sheet：行号 r（表头为 1）对应工单 T{r-1}，与 findings.md 的编号一致；内嵌图片按锚点行归属。
"""

from __future__ import annotations

import argparse
import hashlib
import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import yaml
from PIL import Image

from app.official.models import OFFICIAL_DIR, Dealer, Order, QAEntry, Ticket
from app.settings import REPO_ROOT

SOURCE = REPO_ROOT / "docs" / "【黑客松智能服务赛道】工具资料包" / "黑客松智能服务赛题资料.xlsx"
ORDER_NO = re.compile(r"^Anker\d{3}$")
FULL_DATE = re.compile(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})")
MONTH_DAY = re.compile(r"(\d{1,2})/(\d{1,2})")
QA_SHEETS = {"eufy Robovacs Q&A": "rv", "eufy Breast Pump Q&A": "bp"}
ATTACHMENT_MAX_PX = 1600


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def lines_of(value: Any) -> list[str]:
    return [part.strip() for part in text(value).split("\n") if part.strip()]


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if match := FULL_DATE.search(text(value)):
        return date(*map(int, match.groups()))
    return None


def month_day(raw: str, anchor: date | None) -> date | None:
    """`Delivered on 9/21` 只有月日：取锚点（发货日）所在年份，早于锚点则顺延一年。"""
    if match := FULL_DATE.search(raw):
        return date(*map(int, match.groups()))
    match = MONTH_DAY.search(raw)
    if not match or anchor is None:
        return None
    found = date(anchor.year, int(match.group(1)), int(match.group(2)))
    return found if found >= anchor else found.replace(year=anchor.year + 1)


def tracking(raw: str) -> str:
    return re.sub(r"^Tracking number\d*:\s*", "", raw, flags=re.IGNORECASE).replace(" ", "")


def package(sku: str, carrier: str, tracking_no: str, status_raw: str, ship_raw: Any, order_status: str) -> dict:
    shipped_on = as_date(ship_raw)
    pkg: dict[str, Any] = {
        "tracking_no": tracking(tracking_no) if tracking_no and tracking_no != "/" else None,
        "carrier": carrier if carrier and carrier != "/" else None,
        "skus": [sku],
        "status": "shipped",
        "shipped_on": shipped_on,
        "raw_status": status_raw,
    }
    lowered = status_raw.lower()
    if "未发货" in order_status:
        pkg.update(status="not_shipped", shipped_on=None, expected_ship_on=as_date(ship_raw))
        pkg["raw_status"] = text(ship_raw)
    elif lowered.startswith("delivered"):
        pkg.update(status="delivered", delivered_on=month_day(status_raw, shipped_on))
    elif "return to sender" in lowered:
        pkg["status"] = "returned"
    elif "transit" in lowered or "transist" in lowered:
        pkg.update(status="in_transit", eta=month_day(status_raw.split("at")[-1], shipped_on))
    return pkg


def row_packages(cells: list[Any], sku: str) -> list[dict]:
    carriers, numbers, statuses = lines_of(cells[9]), lines_of(cells[10]), lines_of(cells[11])
    ships = lines_of(cells[12]) if isinstance(cells[12], str) else [cells[12]]
    order_status = text(cells[8])
    count = max(len(numbers), len(statuses), 1 if "未发货" in order_status or carriers else 0)
    result = []
    for i in range(count):
        def pick(items: list, i: int = i) -> Any:
            return items[i] if i < len(items) else (items[-1] if items else "")

        result.append(package(sku, pick(carriers), pick(numbers), pick(statuses), pick(ships), order_status))
    return result


def parse_orders(ws) -> list[dict]:
    orders: list[dict] = []
    for cells in ws.iter_rows(min_row=2, max_col=14, values_only=True):
        sku = text(cells[2])
        if not sku:
            continue
        line = {
            "sku": sku,
            "description": text(cells[3]),
            "qty": int(cells[4] or 1),
            "purchase_date": as_date(cells[5]),
            "warranty_end": as_date(cells[6]),
            "amount": f"{cells[13]:.2f}" if isinstance(cells[13], int | float) else None,
        }
        order_no = text(cells[1])
        if ORDER_NO.match(order_no):
            orders.append({
                "order_no": order_no,
                "purchase_date": line["purchase_date"],
                "address": text(cells[7]),
                "raw_status": text(cells[8]),
                "lines": [line],
                "packages": [],
            })
        current = orders[-1]
        if not ORDER_NO.match(order_no):
            current["lines"].append(line)
        packages = row_packages(list(cells), sku)
        if current["packages"] and not any(p["tracking_no"] or p["carrier"] for p in packages):
            # 子行没有自己的承运商与单号：随主包裹发出。
            current["packages"][0]["skus"].append(sku)
        else:
            current["packages"].extend(packages)
    return [Order.model_validate(o).model_dump(mode="json") for o in orders]


def parse_dealers(ws) -> list[dict]:
    seen: set[tuple[str, str]] = set()
    dealers = []
    for cells in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        country, seller = text(cells[0]), text(cells[1])
        if country and seller and (country, seller) not in seen:
            seen.add((country, seller))
            dealers.append(Dealer(country=country, seller=seller).model_dump())
    return dealers


def parse_qa(ws, prefix: str, sheet: str) -> list[dict]:
    entries = []
    for cells in ws.iter_rows(min_row=2, max_col=10, values_only=True):
        if not text(cells[3]):
            continue
        entries.append(QAEntry(
            id=f"{prefix}-{len(entries) + 1:04d}",
            sheet=sheet,
            directory=[text(c) for c in cells[:3]],
            title=text(cells[3]),
            content=text(cells[4]),
            visible=text(cells[7]) == "Visible",
            pn=[p.strip() for p in text(cells[8]).split(";") if p.strip()],
            pdt=text(cells[9]),
        ).model_dump())
    return entries


def parse_tickets(ws, attachments: dict[int, list[str]]) -> list[dict]:
    tickets = []
    for row, cells in enumerate(ws.iter_rows(min_row=2, max_col=7, values_only=True), start=2):
        if not text(cells[2]):
            continue
        ref = text(cells[3])
        tickets.append(Ticket(
            id=f"T{row - 1}",
            product_line=text(cells[0]),
            product_name=text(cells[1]),
            body=text(cells[2]),
            order_ref=None if ref in ("", "/") else ref,
            order_no=ref if ORDER_NO.match(ref) else None,
            sku=None if text(cells[4]) in ("", "/") else text(cells[4]),
            scenario="shipping" if text(cells[6]).lower().startswith("shipping") else "troubleshooting",
            attachments=attachments.get(row, []),
        ).model_dump())
    return tickets


def export_images(ws, out: Path) -> dict[int, list[str]]:
    out.mkdir(parents=True, exist_ok=True)
    by_row: dict[int, list[str]] = {}
    for img in sorted(ws._images, key=lambda i: i.anchor._from.row):
        row = img.anchor._from.row + 1
        names = by_row.setdefault(row, [])
        name = f"T{row - 1}" + (f"-{len(names) + 1}" if names else "") + ".jpg"
        picture = Image.open(io.BytesIO(img._data())).convert("RGB")
        picture.thumbnail((ATTACHMENT_MAX_PX, ATTACHMENT_MAX_PX))
        picture.save(out / name, "JPEG", quality=85, optimize=True)
        names.append(name)
    return by_row


def dump(path: Path, header: str, payload: dict) -> None:
    body = yaml.safe_dump(payload, allow_unicode=True, sort_keys=False, width=1000)
    path.write_text(header + body, encoding="utf-8")


def build(source: Path, out: Path) -> dict[str, int]:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(source, data_only=True)
    out.mkdir(parents=True, exist_ok=True)
    header = f"# 生成文件，勿手改：python -m app.official.build\n# 来源 {source.name} sha256={digest}\n"
    attachments = export_images(wb["客诉工单信息"], out / "attachments")
    payloads = {
        "orders.yaml": {"orders": parse_orders(wb["订单信息"])},
        "dealers.yaml": {"dealers": parse_dealers(wb["经销商信息"])},
        "qa.yaml": {"entries": [e for s, p in QA_SHEETS.items() for e in parse_qa(wb[s], p, s)]},
        "tickets.yaml": {"tickets": parse_tickets(wb["客诉工单信息"], attachments)},
    }
    for name, payload in payloads.items():
        dump(out / name, header, payload)
    return {name: len(next(iter(p.values()))) for name, p in payloads.items()} | {
        "attachments": sum(len(v) for v in attachments.values())
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="主办方 xlsx → data/official/")
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--out", type=Path, default=OFFICIAL_DIR)
    args = parser.parse_args(argv)
    print(build(args.source, args.out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
